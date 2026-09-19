"""Convert India_BoP_Analysis_v2.xlsx into data/curated.json (the hand-curated analysis layer).

Run:  python scripts/excel_to_json.py [path-to-xlsx]
Every number comes straight from the workbook's cached values; display strings follow the cell's own number format.
"""
import json, re, sys, warnings
from pathlib import Path
import openpyxl

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = Path(r"C:\Users\PC\OneDrive\EA Demo\Final Output\Economic Analysis\output\India_BoP_Analysis_v2.xlsx")

HEADER_FILLS = {"FF1F4E79", "FF2E75B6"}
SECTION_FILLS = {"FF2E75B6", "FFE2EFDA", "FFFFF2CC"}
TOTAL_FILLS = {"FFD9D9D9", "FFD6E4F0", "FF1F3864"}

# sheet -> list of (section-heading prefix, chart spec)
CHARTS = {
    "Summary Dashboard": [("THE BoP WATERFALL", dict(type="bar", title="BoP waterfall, FY2024-25 (USD bn)", label=0, values=[2], names=["Amount"]))],
    "What India Exports": [
        ("SECTION A", dict(type="doughnut", title="Goods exports by category, FY2024-25 (USD bn)", label=0, values=[1], skip="TOTAL")),
        ("SECTION B", dict(type="doughnut", title="Services exports by category, FY2024-25 (USD bn)", label=0, values=[1], skip="TOTAL")),
        ("SECTION C", dict(type="doughnut", title="Remittances by source, FY2024-25 (USD bn)", label=0, values=[1], skip="TOTAL")),
    ],
    "What India Imports": [
        ("SECTION A", dict(type="bar", title="Goods imports by category (USD bn)", label=0, values=[1, 2], names=["FY2024-25", "FY2023-24"], skip="TOTAL")),
        ("SECTION B", dict(type="doughnut", title="Services imports by category, FY2024-25 (USD bn)", label=0, values=[1], skip="TOTAL")),
        ("SECTION C", dict(type="bar", title="Trade balance by country (USD bn)", label=0, values=[3], names=["Trade balance"])),
    ],
    "CAD Evolution (FY20-FY27)": [
        ("SECTION A", dict(type="combo", title="Current account balance (bars) and forex reserves (line), USD bn", label=0, values=[6, 8], names=["Current account balance", "Forex reserves"])),
        ("SECTION A", dict(type="bar", title="Goods exports vs imports (USD bn)", label=0, values=[1, 2], names=["Goods exports", "Goods imports"])),
        ("SECTION C", dict(type="bar", title="Quarterly current account balance (USD bn)", label=0, values=[5], names=["CAB"], only="Q")),
    ],
    "Full BoP Statement": [("__top__", dict(type="bar", title="BoP sub-totals (USD bn)", label=0, values=[1, 2], names=["FY2024-25", "FY2023-24"], only="═", skip="IDENTITY"))],
    "Who Takes Our Assets": [
        ("SECTION A", dict(type="doughnut", title="FDI ownership by source country (cumulative share, %)", label=0, values=[1], skip="TOTAL", pct=True)),
        ("SECTION C", dict(type="bar", title="External debt outstanding by type (USD bn)", label=0, values=[1], parse_str=True)),
    ],
}


def fmt(v, nf):
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        nf = nf or ""
        m = re.search(r"0\.(0+)", nf)
        dec = len(m.group(1)) if m else (0 if nf in ("0", "#,##0") else None)
        if "%" in nf:
            return f"{v*100:.{dec if dec is not None else 1}f}%"
        if dec is None:
            return f"{v:.1f}".rstrip("0").rstrip(".") if abs(v - round(v)) > 1e-9 else f"{int(round(v))}"
        return f"{v:,.{dec}f}"
    return str(v).strip("\r")


def rgb(cell):
    return cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None


def num(s):
    m = re.search(r"-?\d[\d,]*\.?\d*", str(s))
    return float(m.group(0).replace(",", "")) if m else None


def clean_label(s):
    return re.sub(r"^[\s═▶]+", "", str(s)).strip()


def convert_sheet(ws_f, ws_v):
    blocks, sticky_header, table, sec = [], None, None, ""
    specs = list(CHARTS.get(ws_f.title, []))

    def flush():
        nonlocal table
        if table and table["rows"]:
            blocks.append(table)
        table = None

    for r_idx, row in enumerate(ws_f.iter_rows(), start=1):
        cells = [(c, ws_v.cell(row=c.row, column=c.column).value) for c in row]
        nn = [(c, v) for c, v in cells if v is not None]
        if not nn:
            continue
        first, fv = nn[0]
        fill, bold = rgb(first), bool(first.font.b)
        if any(isinstance(v, str) and "\n" in v for _, v in nn):
            flush()
            blocks.append({"t": "cards", "cards": [[x.strip() for x in str(v).split("\n")] for _, v in nn]})
            continue
        if len(nn) == 1:
            text = str(fv).strip()
            is_section = fill in SECTION_FILLS or text.startswith("▶")
            if table is not None and not is_section and r_idx > 2:
                table["rows"].append({"group": True, "cells": [{"d": text}], "cls": "group" if bold else "note"})
                continue
            flush()
            if r_idx == 1:
                blocks.append({"t": "title", "text": text})
            elif r_idx == 2 and not text.startswith("▶"):
                blocks.append({"t": "subtitle", "text": text})
            elif fill in SECTION_FILLS or text.startswith("▶"):
                sec = text
                blocks.append({"t": "section", "text": text.lstrip("▶ ").strip()})
            elif bold:
                blocks.append({"t": "subhead", "text": text})
            else:
                blocks.append({"t": "note", "text": text})
            continue
        if len(nn) >= 3 and fill in HEADER_FILLS:
            flush()
            ncols = max(c.column for c, _ in nn)
            sticky_header = [str(ws_v.cell(row=row[0].row, column=i).value or "") for i in range(1, ncols + 1)]
            table = {"t": "table", "section": sec, "header": sticky_header, "rows": []}
            continue
        if table is None:
            if sticky_header is None:
                continue
            table = {"t": "table", "section": sec, "header": sticky_header, "rows": []}
        ncols = len(table["header"])
        cls = "total" if fill in TOTAL_FILLS else ("sub" if bold else "")
        if len(nn) <= 2 and bold and nn[0][0].column == 1 and cls != "total":
            table["rows"].append({"group": True, "cells": [{"d": str(fv).strip()}, {"d": " ".join(str(v) for _, v in nn[1:])}], "cls": "group"})
            continue
        out = []
        for i in range(1, ncols + 1):
            c = ws_f.cell(row=row[0].row, column=i)
            v = ws_v.cell(row=row[0].row, column=i).value
            cell = {"d": fmt(v, c.number_format)}
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cell["v"] = v
            out.append(cell)
        table["rows"].append({"cells": out, "cls": cls})
    flush()

    result, used = [], set()
    for b in blocks:
        result.append(b)
        if b["t"] == "table":
            for i, (prefix, s) in enumerate(specs):
                if i in used or prefix == "__top__":
                    continue
                if b["section"].lstrip("▶ ").strip().startswith(prefix):
                    ch = build_chart(b, s)
                    if ch:
                        result.append(ch)
                        used.add(i)
    for prefix, s in specs:
        if prefix != "__top__":
            continue
        for j, b in enumerate(result):
            if b["t"] == "table":
                merged = {"header": b["header"], "rows": [r for x in result if x["t"] == "table" for r in x["rows"]]}
                ch = build_chart(merged, s)
                if ch:
                    result.insert(j, ch)
                break
    return result


def build_chart(tbl, s):
    labels, series = [], [[] for _ in s["values"]]
    for r in tbl["rows"]:
        if r.get("group"):
            continue
        raw = r["cells"][s["label"]]["d"]
        lab = clean_label(raw)
        if not lab or (s.get("skip") and s["skip"] in lab.upper()):
            continue
        if s.get("only") and not raw.strip().startswith(s["only"]):
            continue
        vals = []
        for vi in s["values"]:
            c = r["cells"][vi]
            v = c.get("v")
            if v is None and s.get("parse_str"):
                v = num(c["d"])
            vals.append(v)
        if all(v is None for v in vals):
            continue
        if s.get("parse_str") and "TOTAL" in lab.upper():
            continue
        labels.append(lab)
        for k, v in enumerate(vals):
            series[k].append(round(v * 100, 1) if (s.get("pct") and v is not None) else (round(v, 2) if v is not None else None))
    if not labels:
        return None
    names = s.get("names") or [tbl["header"][i] for i in s["values"]]
    return {"t": "chart", "type": s["type"], "title": s["title"], "labels": labels,
            "series": [{"name": n, "data": d} for n, d in zip(names, series)]}


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    wf = openpyxl.load_workbook(src)
    wv = openpyxl.load_workbook(src, data_only=True)
    out = {"source_file": src.name, "sheets": []}
    for ws in wf.worksheets:
        out["sheets"].append({"name": ws.title, "blocks": convert_sheet(ws, wv[ws.title])})
    (ROOT / "data").mkdir(exist_ok=True)
    (ROOT / "data" / "curated.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    for s in out["sheets"]:
        kinds = [b["t"] for b in s["blocks"]]
        print(s["name"], {k: kinds.count(k) for k in sorted(set(kinds))})


if __name__ == "__main__":
    main()
